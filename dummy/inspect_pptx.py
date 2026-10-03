from pptx import Presentation
import os

templates = {
    'winner': r'paradox26\all_certificate_types\07_General_Event_Winner\General_Event_Winner_Template.pptx',
    'participant': r'paradox26\all_certificate_types\08_General_Event_Participant\General_Event_Participant_Template.pptx',
    'guest': r'paradox26\all_certificate_types\06_Guest\Guest_Template.pptx'
}

for k, p in templates.items():
    if not os.path.exists(p):
        print(f"Missing: {p}")
        continue
    prs = Presentation(p)
    slide = prs.slides[0]
    print(f"=== {k.upper()} TEMPLATE ({p}) ===")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p_idx, para in enumerate(shape.text_frame.paragraphs):
                text = para.text.strip()
                if text:
                    print(f'  Shape "{shape.name}": {text}')
